import os
import re
import shutil
import subprocess
import sys
import pandas as pd
import openpyxl

# ==========================================
# FUNZIONI UTILI (SPLITTING E SOSTITUZIONE)
# ==========================================

def split_italian_name(full_name):
    """
    Divide un nominativo completo (Cognome e Nome) in Cognome e Nome separati.
    Gestisce i prefissi comuni italiani (Di, De, Da, Lo, La, Del, Della, Delle, ecc.).
    """
    if not full_name or not isinstance(full_name, str):
        return "", ""
    
    parts = full_name.strip().split()
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    
    prefixes = {"di", "de", "da", "lo", "la", "li", "del", "della", "delle", "delli", "dell", "van", "von", "d'"}
    
    first_word_lower = parts[0].lower().replace("'", "")
    if (first_word_lower in prefixes or parts[0].endswith("'")) and len(parts) > 2:
        cognome = parts[0] + " " + parts[1]
        nome = " ".join(parts[2:])
    else:
        cognome = parts[0]
        nome = " ".join(parts[1:])
        
    return cognome, nome

def find_column(df_cols, candidate_names):
    """
    Trova il nome effettivo della colonna tra i candidati in modo case-insensitive.
    """
    def normalize(s):
        return " ".join(str(s).strip().lower().split())
        
    cols_map = {normalize(c): c for c in df_cols}
    for candidate in candidate_names:
        cand_norm = normalize(candidate)
        if cand_norm in cols_map:
            return cols_map[cand_norm]
    return None

def make_safe_filename(cognome, nome):
    """
    Genera un nome file sicuro per il file system.
    """
    clean_cognome = "".join(c for c in cognome if c.isalnum())
    clean_nome = "".join(c for c in nome if c.isalnum())
    if not clean_cognome and not clean_nome:
        return "ricevuta"
    return f"{clean_cognome}{clean_nome}-ricevuta"

def replace_placeholders(text, mapping):
    """
    Sostituisce i segnaposto tipo {{campo}} in modo case-insensitive.
    """
    if not isinstance(text, str):
        return text
    
    def repl(match):
        key = match.group(1).strip().lower()
        if key == 'città':
            key = 'citta'
            
        if key in mapping:
            return str(mapping[key])
            
        col_mappings = {
            'cognome e nome del bambino': mapping.get('nome_bambino_completo', ''),
            'cognome del bambino': mapping.get('cognome_bambino', ''),
            'nome del bambino': mapping.get('nome_bambino', ''),
            'codice fiscale del bambino': mapping.get('cf_bambino', ''),
            'nominativo genitore a cui intestare la ricevuta': mapping.get('nome_genitore_completo', ''),
            'cognome del genitore a cui intestare la ricevuta': mapping.get('cognome_genitore', ''),
            'nome del genitore a cui intestare la ricevuta': mapping.get('nome_genitore', ''),
            'cognome del genitore': mapping.get('cognome_genitore', ''),
            'nome del genitore': mapping.get('nome_genitore', ''),
            'codice fiscale del genitore a cui intestare la ricevuta': mapping.get('cf_genitore', ''),
            'codice fiscale del genitore': mapping.get('cf_genitore', ''),
            'via di residenza': mapping.get('via', ''),
            'citta': mapping.get('citta', ''),
            'città': mapping.get('citta', ''),
            'cap': mapping.get('cap', '')
        }
        if key in col_mappings:
            return str(col_mappings[key])
            
        return match.group(0)

    return re.sub(r'\{\{?([^{}]+)\}\}?', repl, text)

def find_libreoffice():
    """
    Rileva il percorso dell'eseguibile di LibreOffice su Windows, macOS e Linux.
    """
    import shutil
    
    # 1. Cerca nel PATH di sistema
    for cmd in ["libreoffice", "soffice", "libreoffice.exe", "soffice.exe"]:
        path = shutil.which(cmd)
        if path:
            return path
            
    # 2. Se su Windows, controlla percorsi di installazione standard
    if sys.platform.startswith('win'):
        standard_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"
        ]
        for path in standard_paths:
            if os.path.exists(path):
                return path
                
    # Fallback predefinito
    return "libreoffice"

def convert_xlsx_to_pdf(xlsx_path, pdf_path, output_dir):
    """
    Converte un file XLSX in PDF. Su Windows tenta prima l'uso di Microsoft Excel (via COM),
    altrimenti ripiega su LibreOffice. Su altre piattaforme usa direttamente LibreOffice.
    """
    # 1. Prova con Microsoft Excel via pywin32 su Windows
    if sys.platform.startswith('win'):
        try:
            import win32com.client
            abs_xlsx = os.path.abspath(xlsx_path)
            abs_pdf = os.path.abspath(pdf_path)
            
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            try:
                wb = excel.Workbooks.Open(abs_xlsx)
                # xlTypePDF = 0
                wb.ExportAsFixedFormat(0, abs_pdf)
                wb.Close(False)
                return True
            except Exception as com_err:
                print(f"DEBUG: Conversione nativa con MS Excel fallita: {com_err}")
            finally:
                excel.Quit()
        except ImportError:
            print("DEBUG: Libreria 'pywin32' non installata. Salto la conversione nativa con MS Excel.")
        except Exception as e:
            print(f"DEBUG: Errore generico inizializzazione MS Excel COM: {e}")

    # 2. Fallback su LibreOffice
    libreoffice_bin = find_libreoffice()
    cmd = [
        libreoffice_bin,
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        output_dir,
        xlsx_path
    ]
    
    try:
        result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        if result.returncode == 0:
            return True
        else:
            print(f"DEBUG: Errore LibreOffice (codice {result.returncode}): {result.stderr}")
    except Exception as e:
        print(f"DEBUG: Errore durante l'esecuzione di LibreOffice: {e}")
        
    return False

# ==========================================
# LOGICA PRINCIPALE DI COMPILAZIONE
# ==========================================

def main():
    print("==================================================")
    print("      COMPILATORE AUTOMATICO DI RICEVUTE PDF       ")
    print("==================================================")
    
    # 1. Ricerca dei file nella cartella corrente
    dati_path = None
    template_path = None
    
    current_files = os.listdir('.')
    
    # Cerca file dei paganti
    for f in current_files:
        if 'paganti' in f.lower() and (f.endswith('.csv') or f.endswith('.xlsx') or f.endswith('.xls') or f.endswith('.xlxsx')):
            dati_path = f
            break
            
    # Cerca il modello di ricevuta
    for f in current_files:
        if ('modello' in f.lower() or 'ricevuta' in f.lower()) and f.endswith('.xlsx'):
            template_path = f
            break
            
    # Se non trovati, prova in templates/
    if not dati_path and os.path.exists('templates'):
        for f in os.listdir('templates'):
            if 'paganti' in f.lower() and (f.endswith('.csv') or f.endswith('.xlsx') or f.endswith('.xls')):
                dati_path = os.path.join('templates', f)
                break
                
    if not template_path and os.path.exists('templates'):
        for f in os.listdir('templates'):
            if ('modello' in f.lower() or 'ricevuta' in f.lower()) and f.endswith('.xlsx'):
                template_path = os.path.join('templates', f)
                break

    # Se mancano i file
    if not dati_path:
        print("ERRORE: Impossibile trovare il file dei paganti (es: paganti.csv o paganti.xlsx)")
        print("Assicurati che sia presente in questa cartella.")
        sys.exit(1)
        
    if not template_path:
        print("ERRORE: Impossibile trovare il file Excel modello della ricevuta (es: modello_ricevuta.xlsx)")
        print("Assicurati che sia presente in questa cartella.")
        sys.exit(1)

    print(f"-> File Dati rilevato: {dati_path}")
    print(f"-> File Modello rilevato: {template_path}")
    
    output_dir = "output_ricevute"
    os.makedirs(output_dir, exist_ok=True)
    print(f"-> I PDF verranno salvati in: {os.path.abspath(output_dir)}")
    print("--------------------------------------------------")
    
    try:
        # Carica foglio dati (tutto come testo)
        if dati_path.lower().endswith('.csv'):
            try:
                # Prova prima con la virgola
                df = pd.read_csv(dati_path, dtype=str, sep=',')
                # Se c'è solo una colonna e l'intestazione contiene il punto e virgola, riprova con ;
                if len(df.columns) == 1 and ';' in df.columns[0]:
                    df = pd.read_csv(dati_path, dtype=str, sep=';')
            except Exception:
                # Fallback
                df = pd.read_csv(dati_path, dtype=str, sep=None, engine='python')
        else:
            df = pd.read_excel(dati_path, dtype=str)
            
        df.columns = [str(c).strip() for c in df.columns]
        
        # Gestione speciale se l'Excel contiene dati CSV in un'unica colonna
        if not dati_path.lower().endswith('.csv') and len(df.columns) == 1 and ',' in df.columns[0]:
            import io
            csv_data = df.columns[0] + "\n" + "\n".join(df.iloc[:, 0].dropna())
            df = pd.read_csv(io.StringIO(csv_data), dtype=str)
            df.columns = [str(c).strip() for c in df.columns]
        
        # Riconoscimento flessibile delle colonne
        cols_bambino_cognome = ["cognome del bambino", "cognome bambino"]
        cols_bambino_nome = ["nome del bambino", "nome bambino"]
        cols_bambino_combo = ["cognome e nome del bambino", "nominativo bambino", "nome e cognome del bambino"]
        
        cols_bambino_cf = ["codice fiscale del bambino", "cf bambino", "codice fiscale bambino"]
        
        cols_genitore_cognome = [
            "cognome del genitore a cui intestare la ricevuta",
            "cognome del genitore",
            "cognome genitore"
        ]
        cols_genitore_nome = [
            "nome del genitore a cui intestare la ricevuta",
            "nome del genitore",
            "nome genitore"
        ]
        cols_genitore_combo = [
            "nominativo genitore a cui intestare la ricevuta",
            "nominativo genitore",
            "cognome e nome genitore",
            "cognome e nome del genitore"
        ]
        
        cols_genitore_cf = [
            "codice fiscale del genitore a cui intestare la ricevuta",
            "codice fiscale del genitore",
            "cf genitore",
            "codice fiscale genitore"
        ]
        
        cols_via = ["via di residenza", "via", "indirizzo"]
        cols_citta = ["citta", "città", "città di residenza", "citta di residenza"]
        cols_cap = ["cap"]

        # Trova le colonne effettive
        col_b_cog = find_column(df.columns, cols_bambino_cognome)
        col_b_nom = find_column(df.columns, cols_bambino_nome)
        col_b_combo = find_column(df.columns, cols_bambino_combo)
        
        col_b_cf = find_column(df.columns, cols_bambino_cf)
        
        col_g_cog = find_column(df.columns, cols_genitore_cognome)
        col_g_nom = find_column(df.columns, cols_genitore_nome)
        col_g_combo = find_column(df.columns, cols_genitore_combo)
        
        col_g_cf = find_column(df.columns, cols_genitore_cf)
        
        col_via = find_column(df.columns, cols_via)
        col_citta = find_column(df.columns, cols_citta)
        col_cap = find_column(df.columns, cols_cap)

        missing = []
        if not (col_b_cog and col_b_nom) and not col_b_combo:
            missing.append("Cognome/Nome Bambino")
        if not col_b_cf:
            missing.append("Codice Fiscale del Bambino")
        if not (col_g_cog and col_g_nom) and not col_g_combo:
            missing.append("Cognome/Nome Genitore")
        if not col_g_cf:
            missing.append("Codice Fiscale del Genitore")
        if not col_via:
            missing.append("Via di Residenza")
        if not col_citta:
            missing.append("Città")
        if not col_cap:
            missing.append("CAP")

        if missing:
            print(f"ERRORE: Mancano colonne nell'Excel dei dati:\n{', '.join(missing)}")
            sys.exit(1)
            
        total = len(df)
        print(f"Trovate {total} righe da elaborare. Avvio conversione...")
        
        def safe_str(val):
            if pd.isna(val):
                return ""
            s = str(val).strip()
            if s.endswith(".0"):
                s = s[:-2]
            if s.lower() == 'nan':
                return ""
            return s
            
        for index, row in df.iterrows():
            # Bambino
            if col_b_cog and col_b_nom:
                cognome_bambino = safe_str(row[col_b_cog])
                nome_bambino = safe_str(row[col_b_nom])
                nominativo_bambino = f"{cognome_bambino} {nome_bambino}".strip()
            else:
                nominativo_bambino = safe_str(row[col_b_combo])
                cognome_bambino, nome_bambino = split_italian_name(nominativo_bambino)
                
            cf_bambino = safe_str(row[col_b_cf])
            
            # Genitore
            if col_g_cog and col_g_nom:
                cognome_genitore = safe_str(row[col_g_cog])
                nome_genitore = safe_str(row[col_g_nom])
                nominativo_genitore = f"{cognome_genitore} {nome_genitore}".strip()
            else:
                nominativo_genitore = safe_str(row[col_g_combo])
                cognome_genitore, nome_genitore = split_italian_name(nominativo_genitore)
                
            cf_genitore = safe_str(row[col_g_cf])
            
            via = safe_str(row[col_via])
            citta = safe_str(row[col_citta])
            cap = safe_str(row[col_cap])
            
            row_data = {
                'nome_bambino_completo': nominativo_bambino,
                'cognome_bambino': cognome_bambino,
                'nome_bambino': nome_bambino,
                'cf_bambino': cf_bambino,
                'nome_genitore_completo': nominativo_genitore,
                'cognome_genitore': cognome_genitore,
                'nome_genitore': nome_genitore,
                'cf_genitore': cf_genitore,
                'via': via,
                'citta': citta,
                'cap': cap
            }
            
            filename_base = make_safe_filename(cognome_bambino, nome_bambino)
            pdf_path = os.path.join(output_dir, f"{filename_base}.pdf")
            if os.path.exists(pdf_path):
                print(f"[{index+1}/{total}] Ricevuta già esistente per {nominativo_bambino} ({filename_base}.pdf). Salto.")
                continue

            # Compila Excel
            wb = openpyxl.load_workbook(template_path)
            for ws in wb.worksheets:
                for r in ws.iter_rows():
                    for cell in r:
                        if cell.value is not None and isinstance(cell.value, str):
                            original_value = cell.value
                            new_value = replace_placeholders(cell.value, row_data)
                            if original_value != new_value:
                                cell.value = new_value
                                # Previene il taglio del testo su celle con altezza fissa disabilitando l'a capo
                                # e ridimensionando il font per adattarlo alla larghezza della cella
                                from copy import copy
                                new_alignment = copy(cell.alignment)
                                new_alignment.wrap_text = False
                                new_alignment.shrink_to_fit = True
                                cell.alignment = new_alignment
            
            import tempfile
            temp_xlsx = os.path.join(tempfile.gettempdir(), f"{filename_base}.xlsx")
            wb.save(temp_xlsx)
            wb.close()


            
            # Converte in PDF
            success = convert_xlsx_to_pdf(temp_xlsx, pdf_path, output_dir)
            if not success:
                print(f"[{index+1}/{total}] ERRORE nella conversione PDF per: {nominativo_bambino}")
                if os.path.exists(temp_xlsx):
                    os.remove(temp_xlsx)
                continue

            
            # Pulisce temporaneo
            if os.path.exists(temp_xlsx):
                os.remove(temp_xlsx)
                
            print(f"[{index+1}/{total}] PDF Generato con successo: {filename_base}.pdf")
            
        print("--------------------------------------------------")
        print("PROCESSO COMPLETATO! Trovi i PDF nella cartella 'output_ricevute'.")
        
    except Exception as e:
        print(f"ERRORE GENERALE DI ELABORAZIONE: {str(e)}")
        sys.exit(1)

if __name__ == "__main__":
    main()
